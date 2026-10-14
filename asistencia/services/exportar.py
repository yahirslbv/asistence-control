"""Exportación de una TablaDeReporte a PDF y a Excel (RF-15).

Los nombres de los alumnos vienen de un CSV, así que son datos no confiables:
- Excel: ninguna celda de texto se convierte en fórmula. openpyxl trata como
  fórmula cualquier texto que empiece con «=»; aquí todo texto se guarda como
  texto, y los que empiezan con =, +, - o @ llevan además `quotePrefix`.
- PDF: el texto se escapa (&, < y >) antes de pasarlo a reportlab, que lo
  interpreta como marcado.
"""

import io
import re
import unicodedata
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TIPO_PDF = 'application/pdf'
TIPO_EXCEL = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

# Un texto que empieza así podría interpretarse como fórmula en una hoja de cálculo.
PREFIJOS_DE_FORMULA = ('=', '+', '-', '@', '\t', '\r')

COLOR_ENCABEZADO = '343A40'
COLOR_MARCADA = 'F8D7DA'

AVISO_DE_CARACTERES = (
    'Algunos caracteres no se pueden mostrar en el PDF y aparecen como «?». '
    'El archivo de Excel conserva el texto completo.')


def nombre_de_archivo(base, extension):
    """Nombre de archivo seguro y sin acentos: reporte_ADM-101_2026-2.pdf.

    Solo deja letras, números, punto, guion y guion bajo; lo demás (espacios,
    barras, comillas, saltos de línea) pasa a guion.
    """
    sin_acentos = unicodedata.normalize('NFKD', str(base)).encode('ascii', 'ignore').decode()
    limpio = re.sub(r'[^A-Za-z0-9._-]+', '-', sin_acentos)
    limpio = re.sub(r'-{2,}', '-', limpio).strip('-_.')[:100] or 'reporte'
    return f'{limpio}.{extension}'


# --- Excel -------------------------------------------------------------------

def a_excel(tabla):
    """La tabla como libro de Excel (bytes)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = _titulo_de_hoja(tabla.titulo)

    fila = 1
    _escribir(hoja, fila, 1, tabla.titulo, Font(bold=True, size=14))
    fila += 1
    _escribir(hoja, fila, 1, tabla.subtitulo, Font(color='555555'))
    fila += 2

    fila_del_encabezado = fila
    for columna, nombre in enumerate(tabla.columnas, start=1):
        celda = _escribir(hoja, fila, columna, nombre,
                          Font(bold=True, color='FFFFFF'),
                          relleno=PatternFill('solid', fgColor=COLOR_ENCABEZADO))
        celda.alignment = Alignment(vertical='center', wrap_text=True)
    fila += 1

    ultima = len(tabla.filas) - 1
    for indice, datos in enumerate(tabla.filas):
        es_total = tabla.con_totales and indice == ultima
        relleno = PatternFill('solid', fgColor=COLOR_MARCADA) if indice in tabla.marcadas else None
        for columna, valor in enumerate(datos, start=1):
            celda = _escribir(hoja, fila, columna, valor,
                              Font(bold=True) if es_total else None, relleno=relleno)
            if es_total:
                celda.border = Border(top=Side(style='thin'))
        fila += 1

    fila += 1
    for nota in tabla.notas:
        _escribir(hoja, fila, 1, nota, Font(italic=True, size=9, color='555555'))
        fila += 1

    _ajustar_columnas(hoja, tabla)
    hoja.freeze_panes = f'A{fila_del_encabezado + 1}'
    hoja.page_setup.orientation = 'landscape' if len(tabla.columnas) > 4 else 'portrait'

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _escribir(hoja, fila, columna, valor, fuente=None, relleno=None):
    celda = hoja.cell(row=fila, column=columna)
    if isinstance(valor, int) and not isinstance(valor, bool):
        celda.value = valor                      # los conteos van como números
    else:
        texto = '' if valor is None else str(valor)
        celda.value = texto
        celda.data_type = 's'                    # nunca una fórmula, aunque empiece con «=»
        if texto.startswith(PREFIJOS_DE_FORMULA):
            celda.quotePrefix = True             # y si alguien edita la celda, sigue siendo texto
    if fuente is not None:
        celda.font = fuente
    if relleno is not None:
        celda.fill = relleno
    return celda


def _titulo_de_hoja(titulo):
    """Excel limita el nombre de la hoja a 31 caracteres y prohíbe []:*?/\\ ."""
    limpio = re.sub(r'[\[\]:*?/\\]', ' ', titulo).strip()[:31].strip()
    return limpio or 'Reporte'


def _ajustar_columnas(hoja, tabla):
    for indice, nombre in enumerate(tabla.columnas, start=1):
        largo = max([len(str(nombre))] + [len(str(f[indice - 1])) for f in tabla.filas])
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(largo + 2, 10), 45)


# --- PDF ---------------------------------------------------------------------

def a_pdf(tabla, comprimir=True):
    """La tabla como PDF (bytes). `comprimir=False` deja el texto legible en el
    archivo, solo para las pruebas."""
    pagina = landscape(A4) if len(tabla.columnas) > 4 else A4
    margen = 1.5 * cm
    salida = io.BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=pagina, leftMargin=margen, rightMargin=margen,
        topMargin=margen, bottomMargin=margen + 0.5 * cm,
        title=tabla.titulo, author='Sistema de Control de Asistencia',
        pageCompression=1 if comprimir else 0,
    )

    base = getSampleStyleSheet()['Normal']
    estilo_titulo = ParagraphStyle('Titulo', parent=base, fontName='Helvetica-Bold',
                                   fontSize=15, leading=19, spaceAfter=3)
    estilo_subtitulo = ParagraphStyle('Subtitulo', parent=base, fontSize=9, leading=12,
                                      textColor=colors.HexColor('#555555'), spaceAfter=10)
    estilo_nota = ParagraphStyle('Nota', parent=base, fontSize=8, leading=10,
                                 textColor=colors.HexColor('#555555'), leftIndent=8,
                                 bulletIndent=0)

    sustituidos = [0]

    def limpio(valor):
        """Texto listo para Paragraph: con solo caracteres que la fuente dibuja y
        con &, < y > escapados (Paragraph interpreta marcado)."""
        texto, cambiados = _compatible_con_pdf('' if valor is None else str(valor))
        sustituidos[0] += cambiados
        return escape(texto)

    ancho_util = pagina[0] - 2 * margen
    elementos = [
        Paragraph(limpio(tabla.titulo), estilo_titulo),
        Paragraph(limpio(tabla.subtitulo), estilo_subtitulo),
        _tabla_pdf(tabla, ancho_util, limpio),
        Spacer(1, 8),
    ]
    notas = [limpio(nota) for nota in tabla.notas]
    if sustituidos[0]:
        notas.append(escape(AVISO_DE_CARACTERES))
    elementos += [Paragraph(nota, estilo_nota, bulletText='•') for nota in notas]

    def pie(lienzo, doc):
        lienzo.saveState()
        lienzo.setFont('Helvetica', 8)
        lienzo.setFillColor(colors.HexColor('#777777'))
        lienzo.drawRightString(pagina[0] - margen, margen - 0.4 * cm, f'Página {doc.page}')
        lienzo.restoreState()

    documento.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()


def _compatible_con_pdf(texto):
    """(texto, cuántos caracteres se cambiaron). La fuente estándar del PDF no
    tiene todo Unicode (por ejemplo «ł» o el japonés) y reportlab los omitiría
    en silencio; aquí se cambian por «?» para que el cambio se vea."""
    salida, cambiados = [], 0
    for caracter in texto:
        try:
            caracter.encode('cp1252')
            salida.append(caracter)
        except UnicodeEncodeError:
            salida.append('?')
            cambiados += 1
    return ''.join(salida), cambiados


def _tabla_pdf(tabla, ancho_util, limpio):
    base = getSampleStyleSheet()['Normal']
    numericas = [_es_numerica(tabla, i) for i in range(len(tabla.columnas))]

    def estilo(negrita=False, blanco=False, derecha=False):
        return ParagraphStyle(
            'Celda', parent=base, fontName='Helvetica-Bold' if negrita else 'Helvetica',
            fontSize=8, leading=10, alignment=TA_RIGHT if derecha else TA_LEFT,
            textColor=colors.white if blanco else colors.black)

    def celda(valor, **kwargs):
        return Paragraph(limpio(valor), estilo(**kwargs))

    filas = [[celda(nombre, negrita=True, blanco=True, derecha=numericas[i])
              for i, nombre in enumerate(tabla.columnas)]]
    ultima = len(tabla.filas) - 1
    for indice, datos in enumerate(tabla.filas):
        es_total = tabla.con_totales and indice == ultima
        filas.append([celda(valor, negrita=es_total, derecha=numericas[i])
                      for i, valor in enumerate(datos)])

    estilos = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + COLOR_ENCABEZADO)),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BBBBBB')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]
    for indice in tabla.marcadas:
        estilos.append(('BACKGROUND', (0, indice + 1), (-1, indice + 1),
                        colors.HexColor('#' + COLOR_MARCADA)))
    if tabla.con_totales and tabla.filas:
        estilos.append(('LINEABOVE', (0, ultima + 1), (-1, ultima + 1), 1, colors.black))

    return Table(filas, colWidths=_anchos(tabla, ancho_util), repeatRows=1,
                 style=TableStyle(estilos))


def _es_numerica(tabla, columna):
    """True si la columna son números o porcentajes: se alinea a la derecha."""
    valores = [fila[columna] for fila in tabla.filas]
    return bool(valores) and all(
        isinstance(v, int) or v in ('', '—', 'Sin registros') or re.fullmatch(r'\d+(\.\d+)? %', str(v))
        for v in valores)


def _anchos(tabla, ancho_util):
    """Reparte el ancho de la página según el texto más largo de cada columna."""
    pesos = []
    for i, nombre in enumerate(tabla.columnas):
        largo = max([len(str(nombre))] + [len(str(f[i])) for f in tabla.filas])
        pesos.append(min(max(largo, 6), 40))
    total = sum(pesos)
    return [ancho_util * peso / total for peso in pesos]


# --- Punto de entrada --------------------------------------------------------

FORMATOS = {
    'pdf': (a_pdf, TIPO_PDF),
    'xlsx': (a_excel, TIPO_EXCEL),
}


def exportar(tabla, formato):
    """(contenido, tipo de contenido, nombre de archivo) de la tabla en ese
    formato. Lanza KeyError si el formato no existe."""
    convertir, tipo = FORMATOS[formato]
    return convertir(tabla), tipo, nombre_de_archivo(tabla.nombre_base, formato)
