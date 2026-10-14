"""Pruebas de la exportación de reportes a PDF y Excel (RF-15)."""

import io
import re
import zipfile
from datetime import date

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from asistencia.models import Alumno, Inscripcion
from asistencia.services.exportar import a_excel, a_pdf, nombre_de_archivo
from asistencia.services.reportes import (
    reporte_del_alumno, reporte_del_grupo, tabla_del_alumno, tabla_del_grupo,
)
from asistencia.services.tablas import TablaDeReporte

from .escenarios import EscenarioFijo
from .utils import (
    PASSWORD, RelojFijoMixin, celdas_de_la_tabla, crear_administrador,
    crear_docente, crear_grupo, crear_sesion, inscribir_alumnos,
)

TIPO_EXCEL = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def texto_del_pdf(contenido):
    """Texto de un PDF generado sin comprimir (`a_pdf(..., comprimir=False)`).

    reportlab escribe el texto como (...) Tj, con los acentos en octal
    (WinAnsi) y partiendo el texto en varios fragmentos; aquí se juntan.
    """
    partes = []
    for fragmento in re.finditer(rb'\(((?:\\.|[^\\()])*)\)\s*Tj', contenido):
        crudo = re.sub(rb'\\([0-7]{3})', lambda m: bytes([int(m.group(1), 8)]), fragmento.group(1))
        crudo = crudo.replace(b'\\(', b'(').replace(b'\\)', b')').replace(b'\\\\', b'\\')
        partes.append(crudo.decode('cp1252'))
    return ''.join(partes)


def hoja_de(contenido):
    return load_workbook(io.BytesIO(contenido)).active


def a_texto(valor):
    return '' if valor is None else str(valor)


def filas_de_la_hoja(hoja):
    """Todas las filas de la hoja como listas de texto (los números, como texto)."""
    return [[a_texto(c.value) for c in fila] for fila in hoja.iter_rows()]


class NombreDeArchivoTests(SimpleTestCase):

    def test_el_nombre_normal_se_conserva(self):
        self.assertEqual(
            nombre_de_archivo('reporte_ADM-101_2026-2', 'pdf'), 'reporte_ADM-101_2026-2.pdf')
        self.assertEqual(
            nombre_de_archivo('reporte_ADM-101_2026-2_1000001', 'xlsx'),
            'reporte_ADM-101_2026-2_1000001.xlsx')

    def test_quita_acentos_y_cambia_espacios_y_signos(self):
        self.assertEqual(
            nombre_de_archivo('reporte_ÁDM ñ_2026-2', 'pdf'), 'reporte_ADM-n_2026-2.pdf')

    def test_un_nombre_hostil_no_escapa_de_la_carpeta_ni_rompe_la_cabecera(self):
        for hostil in ('../../etc/passwd', 'a/b\\c', 'A"B\r\nX-Inyectado: y', '<script>.pdf', '   '):
            with self.subTest(hostil=hostil):
                nombre = nombre_de_archivo(hostil, 'pdf')
                self.assertRegex(nombre, r'^[A-Za-z0-9._-]+$')
                self.assertTrue(nombre.endswith('.pdf'))
                self.assertFalse(nombre.startswith('.'))
                self.assertNotIn('..', nombre)

    def test_si_no_queda_nada_usa_un_nombre_generico(self):
        self.assertEqual(nombre_de_archivo('///', 'pdf'), 'reporte.pdf')
        self.assertEqual(nombre_de_archivo('', 'xlsx'), 'reporte.xlsx')

    def test_limita_la_longitud(self):
        self.assertLessEqual(len(nombre_de_archivo('x' * 500, 'pdf')), 104)


class DescargaDelPdfTests(EscenarioFijo):

    def setUp(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)

    def descargar(self, nombre, *args):
        return self.client.get(reverse(nombre, args=args))

    def test_el_pdf_del_grupo_es_un_adjunto_pdf(self):
        respuesta = self.descargar('reporte_grupo_exportar', self.grupo.pk, 'pdf')
        contenido = b''.join(respuesta.streaming_content)

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertEqual(
            respuesta['Content-Disposition'],
            'attachment; filename="reporte_ADM-101_2026-2.pdf"')
        self.assertEqual(respuesta['Cache-Control'], 'private, no-store')
        self.assertTrue(contenido.startswith(b'%PDF'))
        self.assertIn(b'%%EOF', contenido[-20:])

    def test_el_pdf_del_alumno_es_un_adjunto_pdf(self):
        respuesta = self.descargar('reporte_alumno_exportar', self.grupo.pk, self.a.pk, 'pdf')
        contenido = b''.join(respuesta.streaming_content)

        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertEqual(
            respuesta['Content-Disposition'],
            'attachment; filename="reporte_ADM-101_2026-2_1000001.pdf"')
        self.assertTrue(contenido.startswith(b'%PDF'))

    def test_el_pdf_trae_exactamente_las_mismas_cifras_que_la_pantalla(self):
        pantalla = celdas_de_la_tabla(
            self.client.get(reverse('reporte_grupo', args=[self.grupo.pk])).content)
        texto = texto_del_pdf(a_pdf(tabla_del_grupo(reporte_del_grupo(self.grupo)), comprimir=False))

        celdas = pantalla['encabezado'] + [c for f in pantalla['filas'] + pantalla['pie'] for c in f]
        for celda in celdas:
            if celda:
                self.assertIn(celda, texto, f'falta «{celda}» en el PDF')
        for esperado in ('75.0 %', '100.0 %', '87.5 %', 'Sin registros', 'Bajo el mínimo',
                         'Grupo (ponderado)', 'Mínimo de asistencia: 80 %'):
            self.assertIn(esperado, texto)

    def test_el_pdf_del_alumno_trae_el_detalle_de_sus_sesiones(self):
        texto = texto_del_pdf(a_pdf(
            tabla_del_alumno(reporte_del_alumno(self.grupo, self.a)), comprimir=False))

        for esperado in ('12/10/2026', '15/10/2026', 'Presente', 'Retardo', 'Ausente',
                         'Justificado', 'Porcentaje: 75.0 % (3 de 4 registros)'):
            self.assertIn(esperado, texto)

    def test_los_acentos_y_la_enie_se_ven_bien(self):
        Alumno.objects.filter(pk=self.a.pk).update(nombre='Ñandú Pérez García Ü')

        texto = texto_del_pdf(a_pdf(tabla_del_grupo(reporte_del_grupo(self.grupo)), comprimir=False))

        self.assertIn('Ñandú Pérez García Ü', texto)
        for palabra in ('Matrícula', 'Situación', 'Bajo el mínimo', 'Reporte de asistencia · ADM-101',
                        'Administración de Proyectos'):
            self.assertIn(palabra, texto)
        self.assertNotIn('�', texto)

    def test_un_nombre_con_ampersand_y_marcado_no_rompe_el_pdf_y_sale_tal_cual(self):
        nombre = 'Pérez & Hijos <b>negrita</b> <font color="red">x</font>'
        Alumno.objects.filter(pk=self.a.pk).update(nombre=nombre)

        contenido = a_pdf(tabla_del_grupo(reporte_del_grupo(self.grupo)), comprimir=False)
        texto = texto_del_pdf(contenido)

        self.assertTrue(contenido.startswith(b'%PDF'))
        self.assertIn(nombre, texto)        # literal: no se interpretó como marcado

    def test_el_titulo_y_las_notas_tambien_se_escapan(self):
        tabla = TablaDeReporte(
            titulo='A & B <i>', subtitulo='x < y > z', columnas=['N', 'V'],
            filas=[['a', 1]], notas=['nota con <u>marcado</u> & más'])

        texto = texto_del_pdf(a_pdf(tabla, comprimir=False))

        for esperado in ('A & B <i>', 'x < y > z', 'nota con <u>marcado</u> & más'):
            self.assertIn(esperado, texto)

    def test_un_caracter_que_la_fuente_no_dibuja_sale_como_signo_de_interrogacion_con_aviso(self):
        Alumno.objects.filter(pk=self.a.pk).update(nombre='Łukasz 日本')

        texto = texto_del_pdf(a_pdf(tabla_del_grupo(reporte_del_grupo(self.grupo)), comprimir=False))

        self.assertIn('?ukasz ??', texto)
        self.assertIn('Algunos caracteres no se pueden mostrar en el PDF', texto)

    def test_sin_caracteres_especiales_no_aparece_el_aviso(self):
        texto = texto_del_pdf(a_pdf(tabla_del_grupo(reporte_del_grupo(self.grupo)), comprimir=False))

        self.assertNotIn('Algunos caracteres no se pueden mostrar', texto)

    def test_un_grupo_de_muchos_alumnos_ocupa_varias_paginas_sin_error(self):
        grande = crear_grupo(self.docente, clave='GRANDE-1')
        inscribir_alumnos(grande, 120, primera_matricula=8000000)

        contenido = a_pdf(tabla_del_grupo(reporte_del_grupo(grande)))

        self.assertTrue(contenido.startswith(b'%PDF'))
        self.assertGreater(len(re.findall(rb'/Type /Page\b', contenido)), 1)

    def test_un_grupo_sin_alumnos_y_un_alumno_sin_sesiones_se_exportan(self):
        vacio = crear_grupo(self.docente, clave='VACIO-1')
        alumno = inscribir_alumnos(vacio, 1, primera_matricula=9000000)[0]

        for tabla in (tabla_del_grupo(reporte_del_grupo(vacio)),
                      tabla_del_alumno(reporte_del_alumno(vacio, alumno))):
            self.assertTrue(a_pdf(tabla).startswith(b'%PDF'))
            self.assertTrue(a_excel(tabla).startswith(b'PK'))


class DescargaDelExcelTests(EscenarioFijo):

    def setUp(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)

    def bytes_del_grupo(self, grupo=None):
        grupo = grupo or self.grupo
        respuesta = self.client.get(reverse('reporte_grupo_exportar', args=[grupo.pk, 'xlsx']))
        return respuesta, b''.join(respuesta.streaming_content)

    def test_es_un_adjunto_xlsx_que_openpyxl_vuelve_a_abrir(self):
        respuesta, contenido = self.bytes_del_grupo()

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], TIPO_EXCEL)
        self.assertEqual(
            respuesta['Content-Disposition'],
            'attachment; filename="reporte_ADM-101_2026-2.xlsx"')
        self.assertTrue(contenido.startswith(b'PK'))
        self.assertEqual(hoja_de(contenido)['A1'].value, 'Reporte de asistencia · ADM-101')

    def test_las_celdas_coinciden_con_las_de_la_pantalla(self):
        pantalla = celdas_de_la_tabla(
            self.client.get(reverse('reporte_grupo', args=[self.grupo.pk])).content)
        _, contenido = self.bytes_del_grupo()

        filas = filas_de_la_hoja(hoja_de(contenido))
        inicio = next(i for i, f in enumerate(filas) if f[0] == 'Matrícula')
        n = len(pantalla['filas'])

        self.assertEqual(filas[inicio], pantalla['encabezado'])
        self.assertEqual(filas[inicio + 1: inicio + 1 + n], pantalla['filas'])
        self.assertEqual(filas[inicio + 1 + n: inicio + 2 + n], pantalla['pie'])

    def test_las_cifras_van_como_numeros_y_el_porcentaje_como_texto(self):
        _, contenido = self.bytes_del_grupo()
        hoja = hoja_de(contenido)
        fila_de_ana = next(f for f in hoja.iter_rows() if f[1].value == 'Ana Alonso')

        self.assertEqual([c.value for c in fila_de_ana[2:8]], [4, 1, 1, 1, 1, 3])
        self.assertTrue(all(isinstance(c.value, int) for c in fila_de_ana[2:8]))
        self.assertEqual(fila_de_ana[8].value, '75.0 %')
        self.assertEqual(fila_de_ana[9].value, 'Bajo el mínimo')

    def test_el_subtitulo_las_notas_y_el_minimo_estan_en_la_hoja(self):
        _, contenido = self.bytes_del_grupo()
        filas = filas_de_la_hoja(hoja_de(contenido))
        texto = '\n'.join(c for f in filas for c in f)

        self.assertIn('Mínimo de asistencia: 80 %', texto)
        self.assertIn('Los porcentajes se truncan a un decimal.', texto)
        self.assertIn('Alumnos bajo el mínimo: 1.', texto)

    def test_la_fila_bajo_el_minimo_se_resalta_y_la_del_total_va_en_negrita(self):
        _, contenido = self.bytes_del_grupo()
        hoja = hoja_de(contenido)
        ana = next(f for f in hoja.iter_rows() if f[1].value == 'Ana Alonso')
        beto = next(f for f in hoja.iter_rows() if f[1].value == 'Beto Bravo')
        total = next(f for f in hoja.iter_rows() if f[1].value == 'Grupo (ponderado)')

        self.assertEqual(ana[0].fill.fgColor.rgb[-6:], 'F8D7DA')
        self.assertNotEqual(beto[0].fill.fgColor.rgb[-6:], 'F8D7DA')
        self.assertTrue(total[1].font.bold)

    def test_el_excel_del_alumno_coincide_con_la_pantalla(self):
        pantalla = celdas_de_la_tabla(
            self.client.get(reverse('reporte_alumno', args=[self.grupo.pk, self.a.pk])).content)
        respuesta = self.client.get(
            reverse('reporte_alumno_exportar', args=[self.grupo.pk, self.a.pk, 'xlsx']))
        filas = filas_de_la_hoja(hoja_de(b''.join(respuesta.streaming_content)))
        inicio = next(i for i, f in enumerate(filas) if f[0] == 'Fecha')

        self.assertEqual(
            respuesta['Content-Disposition'],
            'attachment; filename="reporte_ADM-101_2026-2_1000001.xlsx"')
        self.assertEqual(filas[inicio], pantalla['encabezado'])
        self.assertEqual(filas[inicio + 1: inicio + 1 + len(pantalla['filas'])], pantalla['filas'])
        self.assertIn('Porcentaje: 75.0 % (3 de 4 registros)', filas[1][0])


class InyeccionDeFormulasEnExcelTests(EscenarioFijo):
    """Los nombres vienen de un CSV: nada que empiece con = + - @ puede ser fórmula."""

    NOMBRES = {
        'a': "=cmd|' /C calc'!A0",
        'b': '+1+1',
        'c': '-2+3',
    }

    def preparar(self):
        for clave, nombre in self.NOMBRES.items():
            Alumno.objects.filter(pk=getattr(self, clave).pk).update(nombre=nombre)
        extra = Alumno.objects.create(matricula='=2+2', nombre='@SUM(1+1)')
        Inscripcion.objects.create(grupo=self.grupo, alumno=extra, fecha_alta=date(2026, 10, 1))
        return extra

    def exportar(self):
        tabla = tabla_del_grupo(reporte_del_grupo(self.grupo))
        return a_excel(tabla)

    def test_ningun_texto_se_convierte_en_formula(self):
        self.preparar()
        contenido = self.exportar()

        hoja = hoja_de(contenido)
        celdas_de_texto = [c for fila in hoja.iter_rows() for c in fila
                           if isinstance(c.value, str) and c.value]
        for celda in celdas_de_texto:
            self.assertEqual(celda.data_type, 's', f'{celda.coordinate} = {celda.value!r}')
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            xml = z.read('xl/worksheets/sheet1.xml').decode('utf-8')
        self.assertNotIn('<f>', xml)
        self.assertNotIn('<f ', xml)

    def test_el_texto_peligroso_queda_intacto_y_marcado_como_texto(self):
        extra = self.preparar()
        hoja = hoja_de(self.exportar())
        por_nombre = {fila[1].value: fila for fila in hoja.iter_rows() if fila[1].value}

        for nombre in (*self.NOMBRES.values(), '@SUM(1+1)'):
            celda = por_nombre[nombre][1]
            self.assertEqual(celda.value, nombre)          # el texto no cambia
            self.assertEqual(celda.data_type, 's')
            self.assertTrue(celda.quotePrefix)
        self.assertEqual(por_nombre['@SUM(1+1)'][0].value, '=2+2')   # la matrícula también
        self.assertEqual(por_nombre['@SUM(1+1)'][0].data_type, 's')
        self.assertTrue(por_nombre['@SUM(1+1)'][0].quotePrefix)
        self.assertEqual(extra.matricula, '=2+2')

    def test_un_nombre_normal_no_lleva_prefijo(self):
        # Sin `preparar()`: los nombres del escenario siguen siendo normales.
        hoja = hoja_de(self.exportar())
        beto = next(f for f in hoja.iter_rows() if f[1].value == 'Beto Bravo')

        self.assertEqual(beto[1].data_type, 's')
        self.assertFalse(beto[1].quotePrefix)

    def test_el_subtitulo_con_una_materia_que_empieza_con_igual_tambien_es_texto(self):
        self.grupo.materia = '=HYPERLINK("http://malo.example","clic")'
        self.grupo.save()

        contenido = self.exportar()
        hoja = hoja_de(contenido)

        self.assertTrue(hoja['A2'].value.startswith('=HYPERLINK'))
        self.assertEqual(hoja['A2'].data_type, 's')
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            self.assertNotIn('<f>', z.read('xl/worksheets/sheet1.xml').decode('utf-8'))

    def test_una_clave_con_signos_prohibidos_en_el_nombre_de_hoja_no_rompe_nada(self):
        self.grupo.clave = 'A/B:1*?[x]'
        self.grupo.save()

        self.client.login(username='docente@asistencia.test', password=PASSWORD)
        respuesta = self.client.get(
            reverse('reporte_grupo_exportar', args=[self.grupo.pk, 'xlsx']))
        hoja = hoja_de(b''.join(respuesta.streaming_content))

        self.assertEqual(respuesta.status_code, 200)
        # Cada signo prohibido pasa a guion: «A/B:1*?[x]» -> «A-B-1-x-».
        self.assertEqual(
            respuesta['Content-Disposition'],
            'attachment; filename="reporte_A-B-1-x-_2026-2.xlsx"')
        self.assertNotRegex(hoja.title, r'[\[\]:*?/\\]')
        self.assertLessEqual(len(hoja.title), 31)


class PermisosDeLaExportacionTests(RelojFijoMixin, TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.docente = crear_docente()
        cls.intruso = crear_docente('intruso@asistencia.test', 'intruso')
        crear_administrador()
        cls.grupo = crear_grupo(cls.docente)
        cls.alumno, cls.baja = inscribir_alumnos(cls.grupo, 2, primera_matricula=7000000)
        Inscripcion.objects.filter(grupo=cls.grupo, alumno=cls.baja).delete()
        crear_sesion(cls.grupo, fecha=date(2026, 10, 12), estado='realizada')

    def urls(self):
        g, a = self.grupo.pk, self.alumno.pk
        return [
            reverse('reporte_grupo_exportar', args=[g, 'pdf']),
            reverse('reporte_grupo_exportar', args=[g, 'xlsx']),
            reverse('reporte_alumno_exportar', args=[g, a, 'pdf']),
            reverse('reporte_alumno_exportar', args=[g, a, 'xlsx']),
        ]

    def test_el_dueno_descarga_los_cuatro_archivos(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)

        for url in self.urls():
            with self.subTest(url=url):
                respuesta = self.client.get(url)
                self.assertEqual(respuesta.status_code, 200)
                self.assertTrue(respuesta['Content-Disposition'].startswith('attachment;'))

    def test_un_grupo_ajeno_responde_404_y_no_entrega_nada(self):
        self.client.login(username='intruso@asistencia.test', password=PASSWORD)

        for url in self.urls():
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 404)

    def test_el_administrador_recibe_403(self):
        self.client.login(username='admin@asistencia.test', password=PASSWORD)

        for url in self.urls():
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_sin_sesion_se_redirige_al_login(self):
        for url in self.urls():
            with self.subTest(url=url):
                self.assertRedirects(
                    self.client.get(url), f"{reverse('login')}?next={url}",
                    fetch_redirect_response=False)

    def test_un_formato_desconocido_responde_404(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)
        g, a = self.grupo.pk, self.alumno.pk

        for formato in ('docx', 'csv', 'html', 'PDF', 'xls'):
            for url in (reverse('reporte_grupo_exportar', args=[g, formato]),
                        reverse('reporte_alumno_exportar', args=[g, a, formato])):
                with self.subTest(url=url):
                    self.assertEqual(self.client.get(url).status_code, 404)

    def test_el_reporte_de_un_alumno_dado_de_baja_no_se_exporta(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)

        for formato in ('pdf', 'xlsx'):
            url = reverse('reporte_alumno_exportar', args=[self.grupo.pk, self.baja.pk, formato])
            with self.subTest(formato=formato):
                self.assertEqual(self.client.get(url).status_code, 404)

    def test_las_pantallas_ofrecen_los_botones_de_descarga(self):
        self.client.login(username='docente@asistencia.test', password=PASSWORD)
        g, a = self.grupo.pk, self.alumno.pk

        grupo = self.client.get(reverse('reporte_grupo', args=[g]))
        alumno = self.client.get(reverse('reporte_alumno', args=[g, a]))

        for respuesta, pdf, xlsx in (
            (grupo, reverse('reporte_grupo_exportar', args=[g, 'pdf']),
             reverse('reporte_grupo_exportar', args=[g, 'xlsx'])),
            (alumno, reverse('reporte_alumno_exportar', args=[g, a, 'pdf']),
             reverse('reporte_alumno_exportar', args=[g, a, 'xlsx'])),
        ):
            self.assertContains(respuesta, f'href="{pdf}"')
            self.assertContains(respuesta, f'href="{xlsx}"')
            self.assertContains(respuesta, 'Descargar PDF')
            self.assertContains(respuesta, 'Descargar Excel')
